"""
Genera la presentazione Credito Salute SQ in formato .pptx professionale.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm

# Palette
BLU_SCURO   = RGBColor(0x1A, 0x3A, 0x5C)
BLU_MEDIO   = RGBColor(0x2E, 0x6D, 0xA8)
VERDE       = RGBColor(0x2D, 0x6A, 0x4F)
GRIGIO_TESTO= RGBColor(0x33, 0x33, 0x33)
BIANCO      = RGBColor(0xFF, 0xFF, 0xFF)
GRIGIO_LIGHT= RGBColor(0xF0, 0xF4, 0xF8)
GRIGIO_RIGA = RGBColor(0xE0, 0xE8, 0xF0)

FONT = "Calibri"

OUTPUT = Path(__file__).parent / "6-presentazione-generale.pptx"

W = Inches(13.33)   # widescreen 16:9
H = Inches(7.5)


def new_prs():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    return prs


def blank_layout(prs):
    return prs.slide_layouts[6]  # completamente vuoto


def fill_solid(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb


def add_rect(slide, left, top, width, height, rgb):
    from pptx.util import Emu
    shape = slide.shapes.add_shape(1, left, top, width, height)  # MSO_SHAPE.RECTANGLE
    shape.line.fill.background()
    fill_solid(shape, rgb)
    return shape


def tf_para(tf, text, size, bold=False, color=BIANCO, align=PP_ALIGN.LEFT, space_before=0):
    p = tf.add_paragraph()
    p.text = text
    p.alignment = align
    p.space_before = Pt(space_before)
    run = p.runs[0] if p.runs else p.add_run()
    run.font.name  = FONT
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    return p


def set_tf(tf, lines, size, bold=False, color=GRIGIO_TESTO, align=PP_ALIGN.LEFT):
    tf.word_wrap = True
    for i, (txt, kwargs) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
            p.text = txt
            p.alignment = align
        else:
            p = tf.add_paragraph()
            p.text = txt
            p.alignment = align
        sz   = kwargs.get("size", size)
        bd   = kwargs.get("bold", bold)
        col  = kwargs.get("color", color)
        sp   = kwargs.get("space_before", 0)
        p.space_before = Pt(sp)
        r = p.runs[0] if p.runs else p.add_run()
        r.font.name  = FONT
        r.font.size  = Pt(sz)
        r.font.bold  = bd
        r.font.color.rgb = col


def add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


# ── helper posizioni ────────────────────────────────────────────────
L  = Cm(1.5)
R  = W - Cm(1.5)
CW = W - Cm(3)      # content width
CT = Cm(3.2)        # content top standard
CB = H - Cm(1.2)    # content bottom


def slide_cover(prs):
    slide = prs.slides.add_slide(blank_layout(prs))
    # sfondo pieno blu scuro
    add_rect(slide, 0, 0, W, H, BLU_SCURO)
    # striscia verde in basso
    add_rect(slide, 0, H - Cm(1.2), W, Cm(1.2), VERDE)
    # titolo
    tb = add_textbox(slide, Cm(2), Cm(2.5), W - Cm(4), Cm(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Credito Salute SQ"
    p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]
    r.font.name = FONT; r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = BIANCO
    # tagline
    tb2 = add_textbox(slide, Cm(2), Cm(5.2), W - Cm(4), Cm(1.8))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "La spesa di tutti i giorni diventa accesso\na cure infermieristiche a domicilio."
    p2.alignment = PP_ALIGN.LEFT
    r2 = p2.runs[0]
    r2.font.name = FONT; r2.font.size = Pt(18); r2.font.bold = False; r2.font.color.rgb = GRIGIO_RIGA
    # firma
    tb3 = add_textbox(slide, Cm(2), H - Cm(2.2), Cm(8), Cm(0.9))
    tf3 = tb3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "Un programma di Salute Quotidiana"
    r3 = p3.runs[0]
    r3.font.name = FONT; r3.font.size = Pt(11); r3.font.bold = False; r3.font.color.rgb = GRIGIO_RIGA


def slide_with_header(prs, title, subtitle=None):
    slide = prs.slides.add_slide(blank_layout(prs))
    # header bar
    add_rect(slide, 0, 0, W, Cm(2.0), BLU_SCURO)
    # striscia verde sinistra
    add_rect(slide, 0, 0, Cm(0.4), H, VERDE)
    # titolo nell'header
    tb = add_textbox(slide, Cm(1.2), Cm(0.15), W - Cm(2), Cm(1.7))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]
    r.font.name = FONT; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = BIANCO
    # sottotitolo opzionale nell'header
    if subtitle:
        tb2 = add_textbox(slide, Cm(1.2), Cm(1.5), W - Cm(2), Cm(0.7))
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        r2 = p2.runs[0]
        r2.font.name = FONT; r2.font.size = Pt(11); r2.font.bold = False; r2.font.color.rgb = GRIGIO_RIGA
    # footer
    add_rect(slide, 0, H - Cm(0.6), W, Cm(0.6), BLU_SCURO)
    tb_f = add_textbox(slide, Cm(1.2), H - Cm(0.56), W - Cm(2), Cm(0.5))
    tf_f = tb_f.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Credito Salute SQ  ·  Salute Quotidiana"
    p_f.alignment = PP_ALIGN.RIGHT
    r_f = p_f.runs[0]
    r_f.font.name = FONT; r_f.font.size = Pt(8); r_f.font.color.rgb = BIANCO
    return slide


def add_body_text(slide, lines, top=CT, left=L, width=CW, size=13):
    tb = add_textbox(slide, left, top, width, H - top - Cm(1.2))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for line in lines:
        if first:
            p = tf.paragraphs[0]; first = False
        else:
            p = tf.add_paragraph()
        text  = line.get("text", "")
        sz    = line.get("size", size)
        bold  = line.get("bold", False)
        color = line.get("color", GRIGIO_TESTO)
        sp    = line.get("space_before", 0)
        align = line.get("align", PP_ALIGN.LEFT)
        indent= line.get("indent", 0)
        p.text = text
        p.alignment = align
        p.space_before = Pt(sp)
        p.level = indent
        if p.runs:
            r = p.runs[0]
        else:
            r = p.add_run()
            r.text = text
        r.font.name  = FONT
        r.font.size  = Pt(sz)
        r.font.bold  = bold
        r.font.color.rgb = color


def add_table_sq(slide, headers, rows, top, left=L, width=CW):
    col_n = len(headers)
    row_n = len(rows) + 1
    tbl = slide.shapes.add_table(row_n, col_n, left, top, width, Cm(0.7) * row_n).table
    # header
    for ci, h in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.fill.solid(); cell.fill.fore_color.rgb = BLU_SCURO
        cell.text = h
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.runs[0] if p.runs else p.add_run()
        r.font.name = FONT; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = BIANCO
    # rows
    for ri, row in enumerate(rows):
        bg = GRIGIO_LIGHT if ri % 2 == 0 else BIANCO
        for ci, val in enumerate(row):
            cell = tbl.cell(ri + 1, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = bg
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            r = p.runs[0] if p.runs else p.add_run()
            r.font.name = FONT; r.font.size = Pt(10); r.font.color.rgb = GRIGIO_TESTO
    return tbl


def slide_2(prs):
    slide = slide_with_header(prs, "Il problema", "Cosa manca oggi")
    add_body_text(slide, [
        {"text": "Accedere a servizi sanitari semplici è più difficile di quanto dovrebbe essere.", "size": 16, "bold": True, "color": BLU_SCURO, "space_before": 8},
        {"text": " ", "size": 8},
        {"text": "Prelievi, medicazioni, iniezioni, controlli di base sono prestazioni semplici.", "size": 13, "space_before": 2},
        {"text": "Eppure organizzarle richiede tempo, spostamenti e spesso costi che molte famiglie rimandano.", "size": 13, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "Nel frattempo, ogni esercizio commerciale spende ogni anno in gadget, volantini e omaggi", "size": 13, "space_before": 2},
        {"text": "che i clienti dimenticano nel giro di qualche giorno.", "size": 13},
        {"text": " ", "size": 8},
        {"text": "C'è un modo per collegare queste due realtà.", "size": 14, "bold": True, "color": VERDE, "space_before": 4},
    ])


def slide_3(prs):
    slide = slide_with_header(prs, "La soluzione", "Il budget promozionale diventa qualcosa che conta")
    add_body_text(slide, [
        {"text": "L'esercizio commerciale destina una quota del proprio budget promozionale a un fondo gestito da Salute Quotidiana.", "size": 13, "space_before": 8},
        {"text": " ", "size": 8},
        {"text": "I clienti accumulano credito con gli acquisti abituali.", "size": 13},
        {"text": "Il credito è spendibile su prestazioni infermieristiche a domicilio — per sé o per un familiare convivente.", "size": 13, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "Nessun gadget. Nessun volantino.", "size": 14, "bold": True, "color": BLU_SCURO, "space_before": 4},
        {"text": "Un beneficio concreto, spendibile, vicino alla vita reale delle persone.", "size": 14, "bold": True, "color": BLU_SCURO},
    ])


def slide_4(prs):
    slide = slide_with_header(prs, "Come funziona", "Tre ruoli. Un meccanismo semplice.")
    top = CT
    col_w = CW / 3
    roles = [
        ("Esercizio commerciale", ["Stanzia il fondo promozionale", "Espone il materiale informativo", "Invita i clienti a partecipare", "Non gestisce nulla di sanitario"], BLU_SCURO),
        ("Cliente", ["Si iscrive all'app", "Carica gli scontrini", "Accumula credito verificato", "Usa il credito per le prestazioni"], BLU_MEDIO),
        ("Salute Quotidiana", ["Gestisce iscrizioni e verifiche", "Gestisce il fondo", "Organizza le prestazioni", "Produce il report finale"], VERDE),
    ]
    for i, (title, items, color) in enumerate(roles):
        lft = L + Cm(0.3) + i * col_w
        # box header
        hdr = add_rect(slide, lft, top, col_w - Cm(0.4), Cm(1.0), color)
        tb_h = add_textbox(slide, lft + Cm(0.1), top + Cm(0.1), col_w - Cm(0.6), Cm(0.8))
        tf_h = tb_h.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = title; p_h.alignment = PP_ALIGN.CENTER
        r_h = p_h.runs[0] if p_h.runs else p_h.add_run()
        r_h.font.name = FONT; r_h.font.size = Pt(12); r_h.font.bold = True; r_h.font.color.rgb = BIANCO
        # box body
        body_top = top + Cm(1.1)
        body_h = H - body_top - Cm(1.0)
        bx = add_rect(slide, lft, body_top, col_w - Cm(0.4), body_h, GRIGIO_LIGHT)
        tb_b = add_textbox(slide, lft + Cm(0.2), body_top + Cm(0.2), col_w - Cm(0.8), body_h - Cm(0.4))
        tf_b = tb_b.text_frame; tf_b.word_wrap = True
        for j, item in enumerate(items):
            pp = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            pp.text = "• " + item
            pp.space_before = Pt(4)
            rr = pp.runs[0] if pp.runs else pp.add_run()
            rr.font.name = FONT; rr.font.size = Pt(11); rr.font.color.rgb = GRIGIO_TESTO


def slide_5(prs):
    slide = slide_with_header(prs, "La formula", "I numeri del programma")
    # left col: formula
    add_body_text(slide, [
        {"text": "15%", "size": 36, "bold": True, "color": BLU_SCURO, "space_before": 8},
        {"text": "della spesa valida del cliente → Credito Salute SQ", "size": 13, "color": GRIGIO_TESTO},
        {"text": " ", "size": 10},
        {"text": "1 euro SQ", "size": 28, "bold": True, "color": VERDE, "space_before": 4},
        {"text": "= 1 euro di sconto sulle prestazioni", "size": 13, "color": GRIGIO_TESTO},
        {"text": " ", "size": 10},
        {"text": "Esempio pratico:", "size": 12, "bold": True, "color": BLU_SCURO, "space_before": 6},
        {"text": "Un cliente che spende 100 euro al mese accumula 15 euro di Credito SQ.", "size": 11},
        {"text": "In due mesi: prelievo + medicazione a costo zero.", "size": 11},
    ], left=L, width=CW * 0.55)
    # right col: parametri
    tb = add_textbox(slide, L + CW * 0.6, CT, CW * 0.38, H - CT - Cm(1.2))
    tf = tb.text_frame; tf.word_wrap = True
    params = [
        ("Fondo tipico", "1.000 euro"),
        ("Clienti per ciclo", "max 20"),
        ("Durata ciclo", "30 giorni"),
        ("Tetto per cliente", "50 euro SQ"),
    ]
    for j, (k, v) in enumerate(params):
        p1 = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p1.text = k; p1.space_before = Pt(10 if j > 0 else 8)
        r1 = p1.runs[0] if p1.runs else p1.add_run()
        r1.font.name = FONT; r1.font.size = Pt(10); r1.font.color.rgb = GRIGIO_TESTO; r1.font.bold = False
        p2 = tf.add_paragraph()
        p2.text = v
        r2 = p2.runs[0] if p2.runs else p2.add_run()
        r2.font.name = FONT; r2.font.size = Pt(16); r2.font.bold = True; r2.font.color.rgb = BLU_SCURO


def slide_6(prs):
    slide = slide_with_header(prs, "Le prestazioni disponibili", "Tutto quello che si può fare con il credito")
    headers = ["Prestazione", "Costo"]
    rows = [
        ("Iniezione intramuscolare I.M. (su prescrizione)", "8 €"),
        ("Prelievo ematico periferico", "10 €"),
        ("Medicazione semplice", "13 €"),
        ("Controllo parametri di base + educazione sanitaria", "16 €"),
        ("Prelievo arterioso", "20 €"),
        ("Gestione medicazione tracheostomia", "30 €"),
        ("Catetere vescicale / cateterismo estemporaneo", "35 €"),
        ("Posizionamento e gestione sondino naso gastrico", "40 €"),
    ]
    add_table_sq(slide, headers, rows, top=CT, width=CW)
    add_body_text(slide, [
        {"text": "Prestazioni programmate · Non urgenti · Erogate a domicilio da infermiere abilitato", "size": 10, "color": VERDE, "align": PP_ALIGN.CENTER},
    ], top=H - Cm(1.6), width=CW)


def slide_7(prs):
    slide = slide_with_header(prs, "Per l'esercizio commerciale", "Cosa cambia per chi aderisce")
    add_body_text(slide, [
        {"text": "Il budget promozionale già esistente produce un effetto che dura.", "size": 14, "bold": True, "color": BLU_SCURO, "space_before": 8},
        {"text": " ", "size": 8},
        {"text": "I clienti ricordano il servizio perché riguarda qualcosa che conta — non un gadget.", "size": 13, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "L'impegno è minimo:", "size": 12, "bold": True, "color": BLU_SCURO, "space_before": 6},
        {"text": "→  Firmare il contratto di adesione", "size": 12, "space_before": 3},
        {"text": "→  Versare il fondo promozionale concordato", "size": 12, "space_before": 2},
        {"text": "→  Invitare i clienti con le proprie parole", "size": 12, "space_before": 2},
        {"text": "→  Ricevere il report aggregato a fine periodo", "size": 12, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "Tutto il resto lo gestisce Salute Quotidiana.", "size": 13, "bold": True, "color": VERDE, "space_before": 4},
    ])


def slide_8(prs):
    slide = slide_with_header(prs, "Per il cliente", "Cosa ottiene chi partecipa")
    add_body_text(slide, [
        {"text": "Credito reale su servizi concreti, accumulato con acquisti già fatti.", "size": 14, "bold": True, "color": BLU_SCURO, "space_before": 8},
        {"text": " ", "size": 8},
        {"text": "Nessun abbonamento. Nessun impegno. Nessuna spesa aggiuntiva.", "size": 13, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "L'infermiere viene a domicilio, su appuntamento, nel giro di pochi giorni.", "size": 13, "space_before": 2},
        {"text": "Senza organizzare nulla di complicato.", "size": 13, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "Il credito si può usare anche per un genitore, un nonno, un coniuge convivente.", "size": 13, "bold": False, "color": GRIGIO_TESTO, "space_before": 2},
        {"text": "Non è solo un vantaggio personale: è qualcosa da mettere al servizio di chi ci sta vicino.", "size": 13, "bold": True, "color": VERDE, "space_before": 2},
    ])


def slide_9(prs):
    slide = slide_with_header(prs, "Perché funziona su scala", "Il modello è replicabile su qualsiasi esercizio locale")
    # griglia esercizi
    esercizi = [
        "Bar / Caffetteria", "Farmacia", "Parafarmacia", "Panificio",
        "Alimentari", "Palestra", "Tabacchi", "Centro estetico",
    ]
    cols = 4
    bw = CW / cols
    bh = Cm(1.1)
    for i, nome in enumerate(esercizi):
        row = i // cols
        col = i % cols
        lft = L + col * bw
        tp  = CT + row * (bh + Cm(0.3))
        color = BLU_SCURO if i % 2 == 0 else BLU_MEDIO
        add_rect(slide, lft, tp, bw - Cm(0.2), bh, color)
        tb = add_textbox(slide, lft + Cm(0.1), tp + Cm(0.1), bw - Cm(0.4), bh - Cm(0.15))
        tf = tb.text_frame
        p = tf.paragraphs[0]; p.text = nome; p.alignment = PP_ALIGN.CENTER
        r = p.runs[0] if p.runs else p.add_run()
        r.font.name = FONT; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BIANCO

    add_body_text(slide, [
        {"text": "Stesso meccanismo. Contesto diverso. Qualità costante.", "size": 13, "bold": True, "color": VERDE, "space_before": 10},
        {"text": "Ogni esercizio porta il proprio fondo e i propri clienti. Salute Quotidiana gestisce tutto.", "size": 12, "space_before": 4},
    ], top=CT + 2 * (bh + Cm(0.3)) + Cm(0.4))


def slide_10(prs):
    slide = slide_with_header(prs, "Il pilot", "Dati reali da un test reale")
    add_body_text(slide, [
        {"text": "Il primo ciclo del programma è un pilot a scala ridotta:", "size": 14, "bold": True, "color": BLU_SCURO, "space_before": 8},
        {"text": "un esercizio · 20 clienti · 30 giorni · fondo definito", "size": 13, "color": BLU_MEDIO, "space_before": 4},
        {"text": " ", "size": 8},
        {"text": "Il pilot produce dati concreti:", "size": 12, "bold": True, "color": BLU_SCURO, "space_before": 6},
        {"text": "quanti si iscrivono, quanto credito accumulano, quanto ne usano, quali prestazioni scelgono.", "size": 12, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "Questi dati sono la base per valutare l'espansione del modello.", "size": 13, "space_before": 2},
        {"text": " ", "size": 8},
        {"text": "Il rischio è contenuto e definito.", "size": 14, "bold": True, "color": VERDE, "space_before": 4},
        {"text": "Il fondo stanziato è il massimo esponibile. Nessun costo variabile aperto.", "size": 12, "space_before": 2},
    ])


def slide_11(prs):
    slide = slide_with_header(prs, "Prossimi passi", "Come aderire o collaborare")
    # due colonne
    col_w = CW / 2 - Cm(0.3)
    for i, (titolo, color, steps) in enumerate([
        ("Esercizio commerciale", BLU_SCURO, [
            "1. Incontro di presentazione con Salute Quotidiana",
            "2. Definizione di fondo e durata del ciclo",
            "3. Firma del contratto di adesione",
            "4. Avvio del programma",
        ]),
        ("Partner / Finanziatore", VERDE, [
            "1. Richiesta risultati del pilot a chiusura",
            "2. Valutazione estensione su uno o più esercizi",
            "3. Definizione accordo di partnership",
            "4. Scalabilità concordata con Salute Quotidiana",
        ]),
    ]):
        lft = L + i * (col_w + Cm(0.6))
        add_rect(slide, lft, CT, col_w, Cm(0.9), color)
        tb_h = add_textbox(slide, lft + Cm(0.2), CT + Cm(0.05), col_w - Cm(0.4), Cm(0.8))
        tf_h = tb_h.text_frame
        p_h = tf_h.paragraphs[0]; p_h.text = titolo; p_h.alignment = PP_ALIGN.CENTER
        r_h = p_h.runs[0] if p_h.runs else p_h.add_run()
        r_h.font.name = FONT; r_h.font.size = Pt(13); r_h.font.bold = True; r_h.font.color.rgb = BIANCO

        tb_b = add_textbox(slide, lft, CT + Cm(1.05), col_w, H - CT - Cm(2.2))
        tf_b = tb_b.text_frame; tf_b.word_wrap = True
        for j, step in enumerate(steps):
            pp = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            pp.text = step; pp.space_before = Pt(8)
            rr = pp.runs[0] if pp.runs else pp.add_run()
            rr.font.name = FONT; rr.font.size = Pt(12); rr.font.color.rgb = GRIGIO_TESTO


def slide_chiusura(prs):
    slide = prs.slides.add_slide(blank_layout(prs))
    add_rect(slide, 0, 0, W, H, BLU_SCURO)
    add_rect(slide, 0, H - Cm(1.2), W, Cm(1.2), VERDE)
    add_rect(slide, 0, 0, Cm(0.4), H, VERDE)

    tb = add_textbox(slide, Cm(2), Cm(2.0), W - Cm(4), Cm(2.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Un modo diverso di fare promozione locale."
    p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]
    r.font.name = FONT; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = BIANCO

    tb2 = add_textbox(slide, Cm(2), Cm(4.3), W - Cm(4), Cm(1.5))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Non serve spendere di più. Serve spendere in modo che le persone ricordino."
    p2.alignment = PP_ALIGN.LEFT
    r2 = p2.runs[0]
    r2.font.name = FONT; r2.font.size = Pt(14); r2.font.bold = False; r2.font.color.rgb = GRIGIO_RIGA

    tb3 = add_textbox(slide, Cm(2), Cm(5.8), Cm(6), Cm(0.8))
    tf3 = tb3.text_frame
    p3 = tf3.paragraphs[0]; p3.text = "Credito Salute SQ · Salute Quotidiana"
    r3 = p3.runs[0]
    r3.font.name = FONT; r3.font.size = Pt(13); r3.font.bold = True; r3.font.color.rgb = VERDE


def main():
    prs = new_prs()
    slide_cover(prs)
    slide_2(prs)
    slide_3(prs)
    slide_4(prs)
    slide_5(prs)
    slide_6(prs)
    slide_7(prs)
    slide_8(prs)
    slide_9(prs)
    slide_10(prs)
    slide_11(prs)
    slide_chiusura(prs)
    prs.save(str(OUTPUT))
    print(f"Salvato: {OUTPUT.name}")


if __name__ == "__main__":
    main()
